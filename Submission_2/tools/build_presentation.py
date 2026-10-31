"""Build eight editable slides and an image of actual saved stdout."""
from pathlib import Path
import json
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = Path(__file__).resolve().parents[1]
M = json.loads((ROOT/'demo/artifacts/metrics.json').read_text())
PAPERS = json.loads((ROOT/'literature/papers.json').read_text())
NAVY, TEAL, INK, GREY, LIGHT, WHITE = '132B43', '087F8C', '203448', '536577', 'F0F4F7', 'FFFFFF'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
prs.core_properties.title = 'Wearable Gesture Controller — ICT 3132 Submission 2'
prs.core_properties.author = 'Arya Choudhary; Sanika Singh; Vishad Sharma'


def color(h): return RGBColor.from_string(h)


def rect(slide, x,y,w,h, fill=LIGHT, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb=color(fill)
    if line: shape.line.color.rgb=color(line)
    else: shape.line.fill.background()
    if radius: shape.adjustments[0] = .08
    return shape


def text(slide, x,y,w,h, content, size=20, bold=False, fill=INK, align=None):
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=shape.text_frame; tf.clear();tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(.02)
    tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='DejaVu Sans';p.font.size=Pt(size)
        p.font.bold=bold;p.font.color.rgb=color(fill)
        p.space_after=Pt(5)
        if align is not None:p.alignment=align
    return shape


def base(title, kicker='ICT 3132 · SUBMISSION 2'):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid();s.background.fill.fore_color.rgb=color(WHITE)
    rect(s,0,0,13.333,.14,TEAL)
    text(s,.65,.35,12,.3,kicker,12,True,TEAL)
    text(s,.65,.9,12,1.05,title,30,True,NAVY)
    rect(s,.65,6.97,12,.012,'D7DFE7')
    text(s,.65,7.08,11.2,.23,'MIT Manipal · V Semester B.Tech. IT · Synthetic software evidence only',10,fill=GREY)
    text(s,12.1,7.04,.5,.3,str(len(prs.slides)),11,fill=GREY,align=PP_ALIGN.RIGHT)
    return s


def note(s, content, refs=()):
    if refs:
        content+='\n\nSources:\n'+'\n'.join(f"{i+1}. {r['title']} ({r['year']}); {r['venue']}; https://doi.org/{r['doi']}; author evidence: {r['source']}" for i,r in enumerate(PAPERS) if r['key'] in refs)
    s.notes_slide.notes_text_frame.text=content


def bullet(s,x,y,w,title,body):
    text(s,x,y,w,.42,title,21,True,NAVY)
    text(s,x,y+.53,w,1.05,body,18,fill=GREY)


# Render the real, saved stdout (not an invented terminal screenshot).
lines=(ROOT/'demo/artifacts/demo_output.txt').read_text().splitlines()
start=next(i for i,l in enumerate(lines) if l.startswith('TRUE CLASS'))
selected=lines[start:start+5]
img=Image.new('RGB',(1940,440),'#132B43');draw=ImageDraw.Draw(img)
def console_font(name, size):
    try:
        return ImageFont.truetype(name, size)
    except OSError:
        return ImageFont.load_default(size=size)
font=console_font('DejaVuSansMono.ttf',34)
boldfont=console_font('DejaVuSansMono-Bold.ttf',31)
draw.text((40,25),'ACTUAL SAVED OUTPUT | SYNTHETIC DATA ONLY',font=boldfont,fill='#8CE0E3')
for i,line in enumerate(selected):draw.text((40,103+57*i),line,font=font,fill='white')
img.save(ROOT/'demo/artifacts/demo_output.png')

s=base('Wearable Gesture Controller', 'ICT 3132 EMBEDDED SYSTEMS · SUBMISSION 2')
text(s,.68,1.83,11.9,1.0,'An IMU-Based Edge Device for\nGesture-Driven Smart Home Control',27,fill=NAVY)
rect(s,.68,3.12,11.95,.67,LIGHT,radius=True)
text(s,.94,3.29,11.4,.35,'CPU-only demonstration • No hardware required • Synthetic data',17,True,TEAL)
team=[('Arya Choudhary','240911438'),('Sanika Singh','240911664'),('Vishad Sharma','240911580')]
for i,(name,reg) in enumerate(team):
    x=.68+i*4.04
    rect(s,x,4.2,3.8,1.4,LIGHT,radius=True)
    text(s,x+.2,4.45,3.4,.4,name,21,True)
    text(s,x+.2,5.02,3.4,.32,reg,18,fill=GREY)
text(s,.7,6.12,11.6,.4,'V Semester B.Tech. IT · School of Computer Engineering · MIT Manipal',16,fill=GREY)
note(s,'Introduce the team and title. Submission 2 adds a small executed software demonstration to the earlier theoretical report. There is no physical wearable or connected appliance. All experimental classification data are synthetic.')

s=base('Problem statement and motivation')
bullet(s,.7,2.0,5.6,'An additional interaction channel','Speech, reach, or fine-motor limitations can make conventional controls difficult to use.')
bullet(s,.7,3.78,5.6,'The engineering problem','Translate a bounded six-channel motion window into a command, while allowing NO_COMMAND.')
rect(s,6.75,1.95,5.85,4.45,LIGHT,radius=True)
text(s,7.02,2.2,5.3,.5,'Why an inertial interface?',23,True,NAVY)
text(s,7.02,3.02,5.2,2.8,'• Acceleration + angular velocity\n• Local temporal recognition\n• No camera or internet dependency\n• Background and uncertain inputs need rejection',21)
text(s,.72,6.2,5.8,.45,'Accessibility benefit requires future user evaluation.',15,True,TEAL)
note(s,'The motivation comes from the supplied synopsis. A gesture interface would complement accessible fallbacks, not serve every disability. Explain that incidental movement and duplicated window outputs make continuous control harder than classifying isolated examples.')

s=base('Literature gap and objectives')
rect(s,.7,1.98,6.0,3.95,LIGHT,radius=True)
text(s,.93,2.18,5.5,.4,'What the recent literature supports',21,True)
text(s,.93,2.88,5.5,2.7,'Wearable models: preserve time and channel structure.\nEmbedded ML: budget activations and kernels, not only weights.\nBenchmarking: separate quality, latency, memory, and energy.',19)
text(s,7.02,2.15,5.5,.45,'Our application-specific gap',22,True,NAVY)
text(s,7.02,2.88,5.45,1.2,'An inspectable six-channel → GRU → reject → command pipeline, with reproducible evidence.',22)
text(s,7.02,4.3,5.4,1.7,'1  Keep recordings separate across splits\n2  Verify preprocessing and prediction\n3  Expose NO_COMMAND\n4  Bound embedded performance claims',18)
text(s,.73,6.22,11.8,.5,'10 recent papers connect wearable recognition, efficient deployment, and embedded evaluation.',15,True,TEAL)
note(s,'Explain the application-specific integration gap: compact temporal recognition, background handling, and a reproducible command decision layer. The literature audit contains the source records, and the project team has confirmed that the references have been checked.',refs=tuple(r['key'] for r in PAPERS))

s=base('Sensor-to-command architecture')
nodes=[('Synthetic input','ax, ay, az\ngx, gy, gz'),('Signal filter','Causal smoothing\nNormalization'),('Window buffer','64 × 6 samples\nStride 32'),('Compact GRU','16 hidden units\n4 output classes'),('Reject + map','Not background\np ≥ 0.80'),('Console output','LIGHT_ON / OFF\nFAN_TOGGLE')]
for i,(head,body) in enumerate(nodes):
    x=.7+i*2.05
    rect(s,x,2.15,1.86,1.65,LIGHT,line=NAVY,radius=True)
    text(s,x+.10,2.32,1.66,.65,head,17,True,NAVY,PP_ALIGN.CENTER)
    text(s,x+.10,3.05,1.66,.65,body,13,fill=INK,align=PP_ALIGN.CENTER)
    if i<5:
        a=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x+1.9),Inches(2.83),Inches(.12),Inches(.25))
        a.fill.solid();a.fill.fore_color.rgb=color(NAVY);a.line.fill.background()
rect(s,.7,4.42,11.95,1.6,LIGHT,radius=True)
text(s,.98,4.62,11.3,.4,'Future physical path — not implemented',21,True,TEAL)
text(s,.98,5.21,11.1,.58,'IMU + MCU sampling → supported inference runtime → optional BLE / Wi-Fi receiver → appliance',19)
text(s,.73,6.34,11.8,.35,'No MCU, board, radio, receiver, or appliance is needed for the current software.',16,fill=GREY)
note(s,'Walk through the solid implemented host pipeline. Six-axis samples are generated in Python, not sampled from real hardware. A future MCU could use an IMU data-ready interrupt or timer and I2C/SPI, with a fixed-size ring buffer. The diagram is made of editable PowerPoint shapes; its matching report figure is TikZ.')

s=base('Methodology and model choice')
cards=[('96 recordings','480 training windows'),('32 recordings','160 validation windows'),('40 recordings','200 test windows')]
for i,(a,b) in enumerate(cards):
    x=.72+i*4.05;rect(s,x,1.98,3.8,1.2,LIGHT,radius=True)
    text(s,x+.15,2.15,3.5,.4,a,24,True,NAVY,PP_ALIGN.CENTER)
    text(s,x+.15,2.69,3.5,.35,b,17,fill=GREY,align=PP_ALIGN.CENTER)
text(s,.76,3.55,5.6,2.5,'Split independent recordings BEFORE windowing.\nFilter → training-only normalization → 64-sample windows.\nSeed 42; five windows per recording.',21)
text(s,6.9,3.55,5.65,2.5,'GRU(6,16) + Linear(16,4): 1,220 parameters.\nAdam; 30 epochs; validation selects checkpoint.\nLSTM comparison is analytical only.',21)
text(s,.76,6.29,11.7,.4,'The generator is deliberately learnable; it is not a validated wrist-motion simulator.',17,True,TEAL)
note(s,'The four labels are background, swipe_left, swipe_right and wave. Synthetic angular-velocity pulses distinguish directions; oscillation defines wave. Frequency, phase, gain and noise vary by independent recording. The complete generator and seed manifest are included. PyTorch GRU uses two bias vectors per gate. Quantization is not executed.')

s=base('Software demonstration and actual results')
s.shapes.add_picture(str(ROOT/'demo/artifacts/demo_output.png'),Inches(.72),Inches(1.98),width=Inches(11.9))
values=[(f"{100*M['accuracy']:.1f}%",'Synthetic test accuracy'),(f"{M['macro_f1']:.3f}",'Synthetic macro-F1'),(f"{M['host_inference_median_ms']:.3f} ms",'Host-only median inference')]
for i,(a,b) in enumerate(values):
    x=.72+i*4.05;rect(s,x,5.0,3.8,1.02,LIGHT,radius=True)
    text(s,x+.1,5.12,3.6,.38,a,25,True,TEAL,PP_ALIGN.CENTER)
    text(s,x+.1,5.65,3.6,.28,b,14,fill=GREY,align=PP_ALIGN.CENTER)
text(s,.76,6.34,11.7,.38,'Perfect toy-data scores verify this generator/pipeline only; real-world accuracy is unknown.',16,True,NAVY)
note(s,'Live demonstration: from Submission_2 run `python demo/demo.py`. The bundled checkpoint loads without retraining. Run `python demo/demo.py --train` to reproduce training and replace metrics. The displayed table is rendered from actual saved stdout, not invented predictions. Classification metrics are before command rejection. Host timing is CPU, batch one, one thread, ten warmups and 100 calls; it excludes acquisition and preprocessing. All 150 gesture windows issued correct commands and no synthetic background window issued a command in the recorded run. This is not a continuous real-world false-activation result.')

s=base('Limitations and future work')
bullet(s,.72,2.0,5.6,'Evidence limits','Synthetic recordings do not cover user variation, tremor, mounting changes, or real incidental motion.')
bullet(s,.72,3.88,5.6,'No hardware claims','MCU latency, peak RAM, flash footprint, power, and battery life have not been measured.')
rect(s,6.88,1.98,5.7,4.53,LIGHT,radius=True)
text(s,7.12,2.2,5.2,.4,'Next steps, if extended',22,True,NAVY)
text(s,7.12,2.93,5.0,3.1,'• Real-data and user evaluation\n• Continuous rejection + cooldown\n• MCU/runtime operator checks\n• Optional 8-bit quantization\n• Optional wireless integration',20)
text(s,.75,6.31,5.85,.48,'Future evaluation: user recordings and on-target profiling.',15,True,TEAL)
note(s,'Explain that the present evaluation uses generated signals. Continuous command handling, user variation, quantization, and target-specific resource measurements form the next stage of development. The command threshold is part of the current software decision layer.')

s=base('Conclusion and reproducible demonstration')
text(s,.76,2.02,11.7,.95,'A small, working software pipeline advances the synopsis without overstating what has been built.',28,True,NAVY)
text(s,.77,3.18,11.55,1.35,'Implemented: six-channel synthetic input, preprocessing, windows, trained GRU, rejection, and command output.\nVerified: held-out synthetic metrics and seven pipeline tests.',21)
rect(s,.76,4.9,11.8,1.12,LIGHT,radius=True)
text(s,.97,5.08,11.3,.68,'Run: python demo/demo.py\nReproduce: python demo/demo.py --train',21,True,TEAL)
text(s,.78,6.28,11.65,.4,'Real-user performance, MCU deployment, quantization, and physical control remain future work.',16,fill=GREY)
note(s,'Finish by running the saved-model demonstration. The package contains reproducible code, recorded simulation metrics, the report, editable diagrams, and literature sources. Present host-only performance separately from prospective MCU measurements.')

# Verify native shape boundaries, then save an editable deck.
for i,s in enumerate(prs.slides,1):
    for sh in s.shapes:
        assert sh.left>=0 and sh.top>=0 and sh.left+sh.width<=prs.slide_width+10 and sh.top+sh.height<=prs.slide_height+10,(i,sh.name)
prs.save(ROOT/'presentation.pptx')
print('Saved eight-slide editable presentation and actual-output image.')
